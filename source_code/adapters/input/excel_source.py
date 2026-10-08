"""Legacy Excel-workbook implementation of the input source port.

FROZEN: the workbook front end is kept for existing decks and for the
regression tests of the two input paths; new input features are YAML-only.
Every read reproduces the pandas/openpyxl call that used to sit in the
model, the loaders or the Simulation facade.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

CONDUCTOR_DEFINITION_FILE = "conductor_definition.xlsx"
CONDUCTOR_SHEET_NAMES = [
    "CONDUCTOR_files",
    "CONDUCTOR_input",
    "CONDUCTOR_operation",
    "CONDUCTOR_coupling",
]


def find_file_containing(directory: Path, substring: str) -> Optional[Path]:
    """The one file of ``directory`` whose name contains ``substring``
    (case-insensitive); None when absent; error when ambiguous."""
    candidates = [
        path for path in directory.iterdir()
        if substring.lower() in path.name.lower()
    ]
    if not candidates:
        return None
    if len(candidates) > 1:
        raise FileExistsError(
            f"Multiple files containing '{substring}' were found in "
            f"{directory}: {', '.join(str(p) for p in candidates)}"
        )
    return candidates[0]


class ExcelInputSource:
    def __init__(self, directory):
        self.directory = Path(directory)
        self.definition_path = self.directory / CONDUCTOR_DEFINITION_FILE
        self.transitory_path = find_file_containing(self.directory, "transitory_input")
        self.grid_path = find_file_containing(self.directory, "grid")
        self.diagnostics_path = find_file_containing(self.directory, "diagnostic")
        self.coupling_path = find_file_containing(self.directory, "coupling")
        workbook = load_workbook(self.definition_path, data_only=True)
        self._definition_sheets = [
            workbook["CONDUCTOR_files"],
            workbook["CONDUCTOR_input"],
            workbook["CONDUCTOR_operation"],
        ]
        self.conductor_sheet_names = [sheet.title for sheet in workbook.worksheets]

    # -- simulation level -------------------------------------------------
    def transient_settings(self) -> Dict[str, Any]:
        settings = pd.read_excel(
            self.transitory_path,
            sheet_name="TRANSIENT",
            skiprows=1,
            header=0,
            index_col=0,
            usecols=["Variable name", "Value"],
        )["Value"].to_dict()
        # Checkpointing settings are YAML-only: force the defaults so both
        # input paths always yield the two keys.
        settings["RESTART"] = False
        settings["AUTOSAVE_INTERVAL"] = None
        return settings

    def environment_settings(self) -> Dict[str, Any]:
        environment_path = find_file_containing(self.directory, "environment")
        return pd.read_excel(
            environment_path,
            sheet_name="ENVIRONMENT",
            header=0,
            index_col=0,
            usecols=["Variable name", "Value"],
        )["Value"].to_dict()

    def hydraulic_network(self) -> Optional[Dict[str, Any]]:
        return None  # YAML-only feature

    @property
    def conductor_count(self) -> int:
        return int(self._definition_sheets[0].cell(row=1, column=2).value)

    def conductor_coupling_dataframe(self):
        return pd.read_excel(
            self.definition_path,
            sheet_name="CONDUCTOR_coupling",
            header=0,
            index_col=0,
        )

    # -- conductor level --------------------------------------------------
    def conductor_name(self, counter: int) -> str:
        return self._definition_sheets[0].cell(row=1, column=1).value

    def conductor_identifier(self, counter: int) -> str:
        return self._definition_sheets[0].cell(row=3, column=4 + counter).value

    def _identifier_sheet(self, sheet_name: str, counter: int) -> Dict[str, Any]:
        identifier = self.conductor_identifier(counter)
        return pd.read_excel(
            self.definition_path,
            sheet_name=sheet_name,
            skiprows=2,
            header=0,
            index_col=0,
            usecols=["Variable name", identifier],
        )[identifier].to_dict()

    def conductor_files(self, counter: int) -> Dict[str, Any]:
        return self._identifier_sheet("CONDUCTOR_files", counter)

    def conductor_inputs(self, counter: int) -> Dict[str, Any]:
        return self._identifier_sheet("CONDUCTOR_input", counter)

    def conductor_operations(self, counter: int) -> Dict[str, Any]:
        return self._identifier_sheet("CONDUCTOR_operation", counter)

    def grid_settings(self, counter: int) -> Dict[str, Any]:
        identifier = self.conductor_identifier(counter)
        return pd.read_excel(
            self.grid_path,
            sheet_name="GRID",
            skiprows=2,
            header=0,
            index_col=0,
            usecols=["Variable name", identifier],
            dtype="object",
        )[identifier].to_dict()

    def _diagnostic_column(self, sheet_name: str, counter: int) -> List[float]:
        identifier = self.conductor_identifier(counter)
        frame = pd.read_excel(
            self.diagnostics_path,
            sheet_name=sheet_name,
            skiprows=2,
            header=0,
            usecols=[identifier],
        )
        return frame.iloc[:, 0].dropna().to_numpy().astype(float)

    def diagnostic_spatial_times(self, counter: int) -> List[float]:
        return self._diagnostic_column("Spatial_distribution", counter)

    def diagnostic_time_positions(self, counter: int) -> List[float]:
        return self._diagnostic_column("Time_evolution", counter)

    def coupling_dataframes(self, counter: int) -> Dict[str, Any]:
        return pd.read_excel(
            self.coupling_path,
            sheet_name=None,
            skiprows=1,
            header=0,
            index_col=0,
        )

    # -- components -------------------------------------------------------
    def _component_workbook_paths(self, counter: int) -> Tuple[Path, Path]:
        files = self.conductor_files(counter)
        return (
            self.directory / files["STRUCTURE_ELEMENTS"],
            self.directory / files["OPERATION"],
        )

    def component_groups(self, counter: int) -> List[Tuple[str, str, List[str]]]:
        input_path, operation_path = self._component_workbook_paths(counter)
        wb_input = load_workbook(input_path, data_only=True)
        wb_operations = load_workbook(operation_path, data_only=True)
        groups = []
        for sheet_name in wb_input.sheetnames:
            sheet = wb_input[sheet_name]
            sheet_operation = wb_operations[sheet_name]
            _check_object_number(input_path, operation_path, sheet, sheet_operation)
            _check_repeated_headings(input_path, sheet)
            _check_repeated_headings(operation_path, sheet_operation)
            _check_headers(input_path, operation_path, sheet, sheet_operation)
            number = int(sheet.cell(row=1, column=2).value)
            identifiers = [
                sheet.cell(row=3, column=4 + ii).value for ii in range(1, number + 1)
            ]
            groups.append((sheet.cell(row=1, column=1).value, sheet.title, identifiers))
        return groups

    def _component_sheet(self, path: Path, sheet_name: str, identifier: str) -> Dict[str, Any]:
        return pd.read_excel(
            path,
            sheet_name=sheet_name,
            skiprows=2,
            header=0,
            index_col=0,
            usecols=["Variable name", identifier],
        )[identifier].to_dict()

    def component_raw_inputs(self, counter, sheet_name, identifier) -> Dict[str, Any]:
        input_path, _ = self._component_workbook_paths(counter)
        return self._component_sheet(input_path, sheet_name, identifier)

    def component_raw_operations(self, counter, sheet_name, identifier) -> Dict[str, Any]:
        _, operation_path = self._component_workbook_paths(counter)
        return self._component_sheet(operation_path, sheet_name, identifier)


# -- workbook consistency checks (formerly conductor.input_validator) --------


def _check_headers(path_input, path_operation, sheet_input, sheet_operation):
    header_input = list(
        pd.read_excel(
            path_input, sheet_name=sheet_input.title, skiprows=2, header=0, index_col=0
        ).columns
    )
    header_operation = list(
        pd.read_excel(
            path_operation,
            sheet_name=sheet_operation.title,
            skiprows=2,
            header=0,
            index_col=0,
        ).columns
    )
    for ii in range(len(header_input)):
        if header_input[ii] != header_operation[ii]:
            raise SyntaxError(
                f"ERROR: object headers in {sheet_input} of file {path_input} and "
                f"in {sheet_operation} of file {path_operation} should be the "
                f"same!\nCompare the third row, column {ii+1} of those sheets.\n"
            )


def _check_object_number(path_1, path_2, sheet_1, sheet_2):
    if int(sheet_1.cell(row=1, column=2).value) != int(
        sheet_2.cell(row=1, column=2).value
    ):
        raise ValueError(
            f"ERROR: number of objects defined in file {path_1} sheet "
            f"{sheet_1.title} and in file {path_2} sheet {sheet_2.title} must be "
            "the same.\nPlease compare cell B1 of those sheets.\n"
        )


def _check_repeated_headings(input_file, sheet):
    columns = list(
        sheet.iter_rows(
            min_row=3,
            max_row=3,
            min_col=sheet.min_column,
            max_col=sheet.max_column,
            values_only=True,
        )
    )[0][4:]
    dict_colum = {column: columns.count(column) for column in columns}
    if max(list(dict_colum.values())) > 1:
        raise ValueError(
            f"ERROR! Different objects of the same kind ({sheet['A1'].value}) can "
            "not have the same identifier.\nUser defines the following:\n"
            f"{dict_colum.items()}.\nPlease check the headers in sheet "
            f"{sheet.title} of file {input_file}"
        )
